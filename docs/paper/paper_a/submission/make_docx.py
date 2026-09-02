"""Flatten the Paper A LaTeX into pandoc-friendly TeX and build the Word
submission set (article, supplementary material, cover letter).

Deterministic, assert-guarded: every substitution that must happen a known
number of times is counted. Cross-references and citation numbers are read
from main.aux / si.aux, never guessed.
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

LATEX = Path(__file__).resolve().parents[1] / 'latex'
SUB = LATEX.parent / 'submission'
BUILD = SUB / 'build'
BUILD.mkdir(exist_ok=True)


def aux_labels(aux_path):
    text = aux_path.read_text(encoding='utf-8')
    out = {}
    for m in re.finditer(r'\\newlabel\{([^}]*)\}\{\{([^{}]*)\}', text):
        out[m.group(1)] = m.group(2)
    return out


def aux_citations(aux_path):
    text = aux_path.read_text(encoding='utf-8')
    order, seen = [], set()
    for m in re.finditer(r'\\citation\{([^}]*)\}', text):
        for k in m.group(1).split(','):
            k = k.strip()
            if k and k not in seen:
                seen.add(k)
                order.append(k)
    return {k: i + 1 for i, k in enumerate(order)}


LABELS = aux_labels(LATEX / 'main.aux')
LABELS_SI = aux_labels(LATEX / 'si.aux')
CITES = aux_citations(LATEX / 'main.aux')
assert len(CITES) == 31, len(CITES)


def clean_colspec(spec):
    spec = re.sub(r'@\{[^{}]*\}', '', spec)
    spec = re.sub(r'>\{[^{}]*\}', '', spec)
    spec = re.sub(r'[pm]\{[^{}]*\}', 'l', spec)
    spec = spec.replace('X', 'l')
    out = ''.join(c for c in spec if c in 'lcr')
    assert out, spec
    return out


def transform(text, labels):
    text = text.replace('\\allowbreak{}', '')
    text = text.replace('\\artifact{', '\\texttt{')
    text = text.replace('\\expid{', '\\texttt{')

    def label_num(key):
        assert key in labels, key
        return labels[key]

    for macro, word in [('sref', 'section'), ('Sref', 'Section'),
                        ('fref', 'figure'), ('Fref', 'Figure'),
                        ('tref', 'table'), ('Tref', 'Table')]:
        text = re.sub(r'\\' + macro + r'\{([^}]*)\}',
                      lambda m, w=word: f'{w}~{label_num(m.group(1))}', text)
    text = re.sub(r'\\ref\{([^}]*)\}', lambda m: label_num(m.group(1)), text)

    def cite_repl(m):
        keys = [k.strip() for k in m.group(2).split(',')]
        for k in keys:
            assert k in CITES, k
        sep = ' ' if m.group(1) else ''
        return sep + '[' + ', '.join(str(CITES[k]) for k in keys) + ']'
    text = re.sub(r'(~)?\\cite\{([^}]*)\}', cite_repl, text)

    # caption numbering: prepend the printed "Figure N."/"Table N." using the
    # label that follows the caption inside the same float.
    def float_repl(m):
        block = m.group(0)
        lab = re.search(r'\\label\{((?:fig|tab):[^}]*)\}', block)
        if lab:
            num = label_num(lab.group(1))
            kind = 'Figure' if lab.group(1).startswith('fig:') else 'Table'
            block = block.replace('\\caption{', f'\\caption{{{kind} {num}. ', 1)
        return block
    text = re.sub(r'\\begin\{(figure|table)\}.*?\\end\{\1\}',
                  float_repl, text, flags=re.DOTALL)
    text = re.sub(r'\\label\{[^}]*\}', '', text)

    SPEC = r'((?:[^{}]|\{[^{}]*\})*)'
    def mc_repl(m):
        keep = ''.join(c for c in re.sub(r'@\{[^{}]*\}', '', m.group(2))
                       if c in 'lcr')
        assert keep, m.group(2)
        return '\\multicolumn{' + m.group(1) + '}{' + keep + '}'
    text = re.sub(r'\\multicolumn\{(\d+)\}\{' + SPEC + r'\}', mc_repl, text)
    text = re.sub(r'\\begin\{tabularx\}\{\\textwidth\}\{' + SPEC + r'\}',
                  lambda m: ('\\begin{tabular}{'
                             + clean_colspec(m.group(1)) + '}'),
                  text)
    text = text.replace('\\end{tabularx}', '\\end{tabular}')
    text = re.sub(r'\\begin\{tabular\}\{' + SPEC + r'\}',
                  lambda m: ('\\begin{tabular}{'
                             + clean_colspec(m.group(1)) + '}'),
                  text)
    text = re.sub(r'\\includegraphics(\[[^]]*\])?\{figures/([^}]*)\.pdf\}',
                  r'\\includegraphics\1{figures/png300/\2.png}', text)
    text = re.sub(r'^\\shorttitle\{[^}]*\}\s*$', '', text, flags=re.MULTILINE)
    return text


def word_friendly_count_table(text):
    """Split the 11-column PDF table into two readable Word tables.

    Values are parsed from the LaTeX table itself and then rearranged, so the
    editable rendition cannot acquire a separately typed numeric source.
    """
    pattern = re.compile(
        r'\\begin\{table\}\[t\]\n'
        r'\\caption\{The four executed mass-ladder rungs:.*?'
        r'\\end\{table\}', re.DOTALL)
    matches = list(pattern.finditer(text))
    assert len(matches) == 1, len(matches)
    block = matches[0].group(0)
    lines = [line.strip() for line in block.splitlines()]
    starts = [i for i, line in enumerate(lines) if line.startswith('$0.')]
    assert len(starts) == 4, starts

    def cells(line):
        parts = [part.strip() for part in line.split('&')]
        assert len(parts) == 11, (len(parts), line)
        parts[-1] = parts[-1].removesuffix(r'\\').strip()
        return parts

    def unmath(value):
        value = value.strip()
        if value.startswith('$') and value.endswith('$'):
            value = value[1:-1]
        return value

    def lower(value):
        value = unmath(value)
        assert value.startswith('[') and value.endswith(','), value
        return value[1:-1]

    def upper(value):
        value = unmath(value).replace(r'\phantom{[}', '')
        value = value.replace(r'{+}', '+')
        assert value.endswith(']'), value
        return value[:-1]

    rows = []
    for i in starts:
        first, second = cells(lines[i]), cells(lines[i + 1])
        verdict = re.fullmatch(r'\\textbf\{([A-Z]+)\}', first[10])
        assert verdict, first[10]
        pilot = unmath(first[3])
        assert pilot.startswith('k='), pilot
        rows.append({
            'mu': unmath(first[0]), 'mass': unmath(first[1]),
            'v_lo': lower(first[2]), 'v_hi': upper(second[2]),
            'pilot': pilot, 'n': first[4], 'k': first[5], 'u': first[6],
            'c_lo': lower(first[7]), 'c_hi': upper(second[7]),
            'd_lo': lower(first[8]), 'd_hi': upper(second[8]),
            'band': unmath(first[9]), 'verdict': verdict.group(1),
        })

    design_rows = '\n'.join(
        f"${r['mu']}$ & ${r['mass']}$ & ${r['pilot']}$ & {r['n']} & "
        f"{r['k']} / {r['u']} & \\textbf{{{r['verdict']}}} \\\\"
        for r in rows)
    interval_rows = '\n'.join(
        f"${r['mu']}$ & certified $V$ & $[{r['v_lo']}, {r['v_hi']}]$ \\\\\n"
        f"${r['mu']}$ & $C$ & $[{r['c_lo']}, {r['c_hi']}]$ \\\\\n"
        f"${r['mu']}$ & $D$ & $[{r['d_lo']}, {r['d_hi']}]$ \\\\\n"
        f"${r['mu']}$ & $B$ & ${r['band']}$ \\\\"
        for r in rows)
    replacement = f"""\\begin{{table}}[t]
\\caption{{The four executed mass-ladder rungs. Part (a) gives the design,
membership counts, and verdict.}}
\\label{{tab:count}}
\\centering
\\small
\\begin{{tabular}}{{@{{}}cccccc@{{}}}}
\\toprule
$\\mu$ & $M$ & pilot & $N$ &
$K_{{\\mathrm{{certain}}}} / U_{{\\mathrm{{amb}}}}$ & verdict \\\\
\\midrule
{design_rows}
\\bottomrule
\\end{{tabular}}
\\end{{table}}

\\begin{{table}}[t]
\\caption{{Table 5 (continued). Part (b) gives the certified-volume and
interval comparison used by the gate.}}
\\centering
\\small
\\begin{{tabular}}{{@{{}}ccc@{{}}}}
\\toprule
$\\mu$ & quantity & value or interval \\\\
\\midrule
{interval_rows}
\\bottomrule
\\end{{tabular}}
\\end{{table}}"""
    return text[:matches[0].start()] + replacement + text[matches[0].end():]


# ---------------------------------------------------------------- article
front = (LATEX / 'sections' / 'front.tex').read_text(encoding='utf-8')
abstract = re.search(r'\\begin\{paperabstract\}\n(.*?)\\end\{paperabstract\}',
                     front, re.DOTALL).group(1).strip()
keywords = re.search(r'\\paperkeywords\{(.*?)\}', front, re.DOTALL).group(1)
keywords = ' '.join(keywords.split())

order = ['s01_intro', 's02_ladder', 's03_methods', 's04_results',
         's05_negative', 's06_capstone', 's067_schwarzschild', 's068_count',
         's069_oracle', 's07_discussion', 's08_claims', 's10_repro',
         'backmatter']
body = '\n\n'.join(
    (LATEX / 'sections' / f'{name}.tex').read_text(encoding='utf-8')
    for name in order)
refs = (LATEX / 'sections' / 'references.tex').read_text(encoding='utf-8')
refs = refs.replace('\\begin{thebibliography}{99}', '\\section*{References}')
refs = refs.replace('\\end{thebibliography}', '')
n_bib = 0


def bib_repl(m):
    global n_bib
    n_bib += 1
    key = m.group(1)
    assert CITES[key] == n_bib, (key, n_bib)
    return f'[{n_bib}]'


refs = re.sub(r'\\bibitem\{([^}]*)\}', bib_repl, refs)
assert n_bib == 31

TITLE = ('Spacetime quantities from causal order: an operational ladder '
         'with preregistered validations at finite density')
article = f"""\\title{{{TITLE}}}
\\author{{Juneyoung Kim \\\\ Independent researcher \\\\ lpaiu.cs@gmail.com}}
\\date{{}}
\\maketitle

\\section*{{Abstract}}
{abstract}

\\noindent\\textbf{{Keywords:}} {keywords}

\\noindent Submitted to: \\emph{{Classical and Quantum Gravity}}

{body}

{refs}
"""
article = word_friendly_count_table(article)
article = transform(article, LABELS)
(BUILD / 'article_flat.tex').write_text(article, encoding='utf-8')

# ----------------------------------------------------------------- SI doc
si_body = '\n\n'.join(
    (LATEX / 'sections' / f'{name}.tex').read_text(encoding='utf-8')
    for name in ['appendix_a', 'appendix_b'])
# S-numbered headings become literal (pandoc numbering would restart at 1)
si_body = si_body.replace(
    '\\section{Conventions and normalizations}',
    '\\section*{S1. Conventions and normalizations}', 1)
si_body = si_body.replace(
    '\\section{Capstone execution provenance}',
    '\\section*{S2. Capstone execution provenance}', 1)
si_labels = dict(LABELS)
si_labels.update(LABELS_SI)
si = f"""\\title{{Supplementary material for “{TITLE}”}}
\\author{{Juneyoung Kim \\\\ Independent researcher \\\\ lpaiu.cs@gmail.com}}
\\date{{}}
\\maketitle

\\noindent This document carries the article's conventions table
(section~S1) and the execution-provenance record of the preregistered
stages (section~S2). Section numbers without an S prefix refer to the
article.

{si_body}
"""
si = transform(si, si_labels)
(BUILD / 'si_flat.tex').write_text(si, encoding='utf-8')

# ------------------------------------------------------------------ build
SUB.mkdir(exist_ok=True)


def pandoc(args):
    r = subprocess.run(['pandoc'] + args, cwd=LATEX, capture_output=True,
                       text=True)
    sys.stdout.write(r.stdout)
    sys.stderr.write(r.stderr)
    assert r.returncode == 0, r.returncode


def polish_cover_letter(path):
    """Apply restrained journal-correspondence typography after pandoc."""
    doc = Document(path)
    for section in doc.sections:
        section.top_margin = Inches(0.70)
        section.bottom_margin = Inches(0.70)
        section.left_margin = Inches(0.90)
        section.right_margin = Inches(0.90)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.0
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name == 'Captioned Figure':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
        elif paragraph.style.name == 'Image Caption':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(8.5)
                run.font.italic = True
    doc.save(path)


def polish_article(path):
    """Keep the dense mass-ladder table legible in the portrait Word file."""
    doc = Document(path)
    assert len(doc.tables) == 7, len(doc.tables)
    tables = doc.tables[4:6]
    assert [len(table.columns) for table in tables] == [6, 3]
    assert tables[0].cell(0, 5).text == 'verdict'

    style_name = 'CQG Small Table'
    if style_name in doc.styles:
        small = doc.styles[style_name]
    else:
        small = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    small.font.name = 'Times New Roman'
    small.font.size = Pt(8.5)
    small.paragraph_format.space_before = Pt(0)
    small.paragraph_format.space_after = Pt(0)
    small.paragraph_format.line_spacing = 1.0

    widths_by_table = (
        (0.65, 0.55, 0.75, 1.20, 1.25, 1.25),
        (0.75, 1.15, 3.75),
    )
    for table, widths in zip(tables, widths_by_table, strict=True):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.first_child_found_in('w:tblLayout')
        if layout is None:
            layout = OxmlElement('w:tblLayout')
            tbl_pr.append(layout)
        layout.set(qn('w:type'), 'fixed')
        for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst,
                                   widths, strict=True):
            grid_col.set(qn('w:w'), str(round(width * 1440)))

        for row in table.rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.first_child_found_in('w:tcW')
                if tc_w is None:
                    tc_w = OxmlElement('w:tcW')
                    tc_pr.append(tc_w)
                tc_w.set(qn('w:type'), 'dxa')
                tc_w.set(qn('w:w'), str(round(width * 1440)))
                tc_mar = tc_pr.first_child_found_in('w:tcMar')
                if tc_mar is None:
                    tc_mar = OxmlElement('w:tcMar')
                    tc_pr.append(tc_mar)
                for edge in ('top', 'left', 'bottom', 'right'):
                    margin = tc_mar.find(qn(f'w:{edge}'))
                    if margin is None:
                        margin = OxmlElement(f'w:{edge}')
                        tc_mar.append(margin)
                    margin.set(qn('w:w'), '45')
                    margin.set(qn('w:type'), 'dxa')
                for paragraph in cell.paragraphs:
                    paragraph.style = small
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.keep_together = True
                    for run in paragraph.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(8.5)
                # Office Math runs need an explicit size in LibreOffice.
                for math_run in cell._tc.iter(qn('m:r')):
                    run_pr = math_run.find(qn('w:rPr'))
                    if run_pr is None:
                        run_pr = OxmlElement('w:rPr')
                        math_pr = math_run.find(qn('m:rPr'))
                        math_run.insert(1 if math_pr is not None else 0,
                                        run_pr)
                    for tag in ('w:sz', 'w:szCs'):
                        size = run_pr.find(qn(tag))
                        if size is None:
                            size = OxmlElement(tag)
                            run_pr.append(size)
                        size.set(qn('w:val'), '17')
    doc.save(path)


options = set(sys.argv[1:])
assert options <= {'--cover-only'}, options
if '--cover-only' not in options:
    pandoc(['-f', 'latex', '-t', 'docx', '--number-sections',
            str(BUILD / 'article_flat.tex'),
            '-o', str(SUB / 'paper_a_article.docx')])
    polish_article(SUB / 'paper_a_article.docx')
    pandoc(['-f', 'latex', '-t', 'docx',
            str(BUILD / 'si_flat.tex'),
            '-o', str(SUB / 'paper_a_supplementary.docx')])
pandoc(['-f', 'markdown', '-t', 'docx',
        str(SUB / 'cover_letter.md'),
        '-o', str(SUB / 'cover_letter.docx')])
polish_cover_letter(SUB / 'cover_letter.docx')
print('docx set written to', SUB)
